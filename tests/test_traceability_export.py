import copy
import csv
import json
import os
import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from core.design_context import construir_contexto_diseno
from core.design_traceability import (
    construir_filas_matriz_diseno,
    evolucionar_matriz_a_diseno,
    normalizar_matriz_requerimientos,
    resumir_trazabilidad_diseno,
)
from core.traceability_export import exportar_csv_excel, exportar_filas_xlsx
from core.utils import TRACEABILITY_COLUMNS, construir_filas_trazabilidad
from database.repository import DB_PATH
from integrations.issue_service import obtener_issues_diseno
from tests.fixtures_diseno import EVALUACION_POR_DISENO, RESULTADOS_CENTRAL_DISENO


PROJECT_ID = os.getenv("GITLAB_PROJECT_ID")

MATRIZ_PATH = Path(__file__).parent.parent / "output" / "csv" / "matriz_trazabilidad_5_HU.csv"
OUTPUT_CSV_DIR = Path(__file__).parent.parent / "output" / "csv"
OUTPUT_XLSX_DIR = Path(__file__).parent.parent / "output" / "xlsx"
OUTPUT_XLSX_DIR.mkdir(parents=True, exist_ok=True)


# ==========================================================
# PASO 7 — Matriz de Requerimientos (datos reales de cache_results)
# ==========================================================

conn = sqlite3.connect(DB_PATH)
cursor = conn.cursor()
cursor.execute("SELECT central_json FROM cache_results")
filas_cache = cursor.fetchall()
conn.close()

batch_results_por_iid = {}
for (raw_central,) in filas_cache:
    central = json.loads(raw_central)
    iid = central.get("issue_iid")
    if iid is None:
        continue
    batch_results_por_iid[iid] = {"status": "ok", "issue_iid": iid, "central": central}

batch_results = list(batch_results_por_iid.values())
assert batch_results, "No hay datos reales en cache_results para probar la matriz de Requerimientos."

filas_requerimientos = construir_filas_trazabilidad(batch_results)

csv_requerimientos_path = OUTPUT_CSV_DIR / "matriz_requerimientos.csv"
xlsx_requerimientos_path = OUTPUT_XLSX_DIR / "matriz_requerimientos.xlsx"

exportar_csv_excel(filas_requerimientos, csv_requerimientos_path)
exportar_filas_xlsx(filas_requerimientos, xlsx_requerimientos_path, nombre_hoja="Trazabilidad Requerimientos")

print("\n" + "=" * 70)
print("MATRIZ DE REQUERIMIENTOS — export")
print("=" * 70)
print("CSV:", csv_requerimientos_path)
print("XLSX:", xlsx_requerimientos_path)
print("Filas:", len(filas_requerimientos))

# ✓ columnas separadas correctamente (delimiter ';')
with csv_requerimientos_path.open(encoding="utf-8-sig") as handle:
    primera_linea = handle.readline()
assert ";" in primera_linea

with csv_requerimientos_path.open(encoding="utf-8-sig", newline="") as handle:
    lector = csv.DictReader(handle, delimiter=";")
    filas_leidas_req = list(lector)

assert filas_leidas_req
for row in filas_leidas_req:
    assert "Código" in row
    assert "Nombre" in row
    assert "Descripción" in row

# ✓ UTF-8 / acentos
with csv_requerimientos_path.open(encoding="utf-8-sig") as handle:
    contenido_req = handle.read()
assert "Código" in contenido_req
assert "Descripción" in contenido_req
assert "Fecha de generación" in contenido_req

# ✓ XLSX: encabezados en celdas separadas, no una cadena CSV completa
wb_req = load_workbook(xlsx_requerimientos_path)
ws_req = wb_req["Trazabilidad Requerimientos"]
for indice, columna in enumerate(TRACEABILITY_COLUMNS, start=1):
    assert ws_req.cell(row=1, column=indice).value == columna
primer_codigo = ws_req.cell(row=2, column=1).value
assert primer_codigo in {row["Código"] for row in filas_leidas_req}
assert ";" not in str(primer_codigo)


# ==========================================================
# PASO 8 — Matriz de Diseño
# ==========================================================

with MATRIZ_PATH.open(encoding="utf-8-sig", newline="") as handle:
    matriz_original = list(csv.DictReader(handle))

matriz_original_copia = copy.deepcopy(matriz_original)
matriz_original_bytes = MATRIZ_PATH.read_bytes()

issues_diseno = obtener_issues_diseno(PROJECT_ID, milestone_title="Diseño")
issues_por_diseno_id = {issue.get("diseno_id"): issue for issue in issues_diseno}
contextos_por_diseno_id = {
    diseno_id: construir_contexto_diseno(issue, matriz_original)
    for diseno_id, issue in issues_por_diseno_id.items()
}


# Resultados de Design_Central / design_summary: fixtures con datos reales ya
# validados (ver tests/fixtures_diseno.py) — sin invocar LLM en esta prueba
# de exportación. La validación real contra Groq/NVIDIA vive en
# tests/test_design_full_flow.py.
resultados_central_diseno = RESULTADOS_CENTRAL_DISENO
evaluacion_por_diseno = EVALUACION_POR_DISENO

# DIS-003 sigue bloqueado: nunca participa
contexto_dis003 = contextos_por_diseno_id["DIS-003"]
assert contexto_dis003["validacion_trazabilidad"]["referencias_invalidas"]

matriz_normalizada = normalizar_matriz_requerimientos(matriz_original)
estructura_interna = evolucionar_matriz_a_diseno(matriz_normalizada, resultados_central_diseno)
filas_diseno = construir_filas_matriz_diseno(estructura_interna, evaluacion_por_diseno)
resumen_diseno = resumir_trazabilidad_diseno(filas_diseno)

csv_diseno_path = OUTPUT_CSV_DIR / "matriz_diseno.csv"
xlsx_diseno_path = OUTPUT_XLSX_DIR / "matriz_diseno.xlsx"

exportar_csv_excel(filas_diseno, csv_diseno_path)
exportar_filas_xlsx(
    filas_diseno, xlsx_diseno_path,
    nombre_hoja="Trazabilidad Diseño", resumen=resumen_diseno,
)

print("\n" + "=" * 70)
print("MATRIZ DE DISEÑO — export")
print("=" * 70)
print("CSV:", csv_diseno_path)
print("XLSX:", xlsx_diseno_path)
print("Filas:", len(filas_diseno))
print("Resumen:", resumen_diseno)

# ✓ matriz original (fuente CSV de Requerimientos) intacta
assert matriz_original == matriz_original_copia
assert MATRIZ_PATH.read_bytes() == matriz_original_bytes

# ✓ columnas separadas correctamente (delimiter ';')
with csv_diseno_path.open(encoding="utf-8-sig") as handle:
    primera_linea_diseno = handle.readline()
assert ";" in primera_linea_diseno

with csv_diseno_path.open(encoding="utf-8-sig", newline="") as handle:
    lector_diseno = csv.DictReader(handle, delimiter=";")
    filas_leidas_diseno = list(lector_diseno)

assert filas_leidas_diseno
for row in filas_leidas_diseno:
    assert "Código requisito" in row
    assert "Nombre del requisito" in row
    assert "Descripción" in row
    assert "Estado de trazabilidad" in row
    assert "Estado de Diseño" in row

# ✓ UTF-8 / acentos
with csv_diseno_path.open(encoding="utf-8-sig") as handle:
    contenido_diseno = handle.read()
assert "Diseño" in contenido_diseno
assert "Código requisito" in contenido_diseno
assert "Descripción" in contenido_diseno

# ✓ DIS-003 ausente del CSV de Diseño
assert "DIS-003" not in contenido_diseno
assert "RF-009" not in contenido_diseno
assert "RF-010" not in contenido_diseno

# ✓ XLSX de Diseño: encabezados exactos en celdas separadas
wb_diseno = load_workbook(xlsx_diseno_path)
ws_diseno = wb_diseno["Trazabilidad Diseño"]
assert ws_diseno["A1"].value == "HU origen"
assert ws_diseno["B1"].value == "Código requisito"
columnas_diseno = list(filas_diseno[0].keys())
for indice, columna in enumerate(columnas_diseno, start=1):
    assert ws_diseno.cell(row=1, column=indice).value == columna

valores_columna_b = {ws_diseno.cell(row=fila, column=2).value for fila in range(2, ws_diseno.max_row + 1)}
assert "DIS-003" not in valores_columna_b
for valor in valores_columna_b:
    assert ";" not in str(valor)

assert "Resumen" in wb_diseno.sheetnames
ws_resumen = wb_diseno["Resumen"]
resumen_leido = {ws_resumen.cell(row=fila, column=1).value: ws_resumen.cell(row=fila, column=2).value for fila in range(1, ws_resumen.max_row + 1)}
assert resumen_leido["cubiertos"] == resumen_diseno["cubiertos"]

print("\nTodas las verificaciones de exportación (CSV/XLSX) pasaron.")
