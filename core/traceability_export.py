import csv

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

COLUMNAS_MULTILINEA = {"Descripción", "Observación"}


def exportar_csv_excel(filas, ruta):
    """
    Exporta una lista plana de diccionarios a CSV legible en Excel: cada
    campo en su propia columna (delimiter=';'), con BOM UTF-8 para que los
    acentos se muestren correctamente. Reutilizable para Requerimientos,
    Diseño, Codificación y Pruebas.
    """
    if not filas:
        return ruta

    columnas = list(filas[0].keys())

    with open(
        ruta,
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as archivo:
        writer = csv.DictWriter(
            archivo,
            fieldnames=columnas,
            delimiter=";",
            quoting=csv.QUOTE_MINIMAL,
        )

        writer.writeheader()
        writer.writerows(filas)

    return ruta


def exportar_filas_xlsx(filas, ruta, nombre_hoja="Trazabilidad", resumen=None):
    """
    Exporta una lista plana de diccionarios a un libro XLSX nativo (openpyxl,
    sin LibreOffice): encabezados en la primera fila, cada campo en su celda,
    fila 1 congelada, filtros activos, columnas ajustadas y salto de línea
    habilitado en Descripción/Observación. Si se recibe `resumen`, agrega una
    segunda hoja "Resumen" con sus pares clave/valor.
    """
    if not filas:
        return ruta

    columnas = list(filas[0].keys())

    workbook = Workbook()
    hoja = workbook.active
    hoja.title = nombre_hoja

    hoja.append(columnas)
    for fila in filas:
        hoja.append([fila.get(columna, "") for columna in columnas])

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    for indice, columna in enumerate(columnas, start=1):
        letra = get_column_letter(indice)
        longitud = max(
            [len(str(columna))] + [len(str(fila.get(columna, ""))) for fila in filas]
        )
        hoja.column_dimensions[letra].width = min(longitud + 2, 60)

        if columna in COLUMNAS_MULTILINEA:
            for celda in hoja[letra][1:]:
                celda.alignment = celda.alignment.copy(wrap_text=True)

    _estilizar_hoja_matriz(hoja, columnas, len(filas))

    if resumen:
        hoja_resumen = workbook.create_sheet("Resumen")
        for clave, valor in resumen.items():
            hoja_resumen.append([clave, valor])
        hoja_resumen.column_dimensions["A"].width = 30
        hoja_resumen.column_dimensions["B"].width = 15
        _estilizar_hoja_resumen(hoja_resumen)

    workbook.save(ruta)
    _marcar_estilos_aplicados(ruta)
    return ruta


_NAVY = "FF1F3864"
_ZEBRA = "FFF5F8FC"
_BORDE = "FFBFC8D6"
_LIGHT = "FFEAF0F8"
_VERDE = "FFE2F0D9"
_AMBAR = "FFFFF2CC"
_ROJO = "FFF8D7DA"

# Colores de estado por palabra clave (solo apariencia; el valor no cambia).
_ESTADOS_ROJO = ("no cubierto", "error", "fallid", "rechaz", "no implementado", "no evaluable")
_ESTADOS_AMBAR = ("pendiente", "declarado", "corregir", "revis", "sin confirmar", "parcial")
_ESTADOS_VERDE = ("conforme", "aprobad", "cubierto", "implementado", "verificad", "completo")


def _color_estado(valor):
    texto = str(valor or "").strip().lower()
    if not texto:
        return None
    if any(k in texto for k in _ESTADOS_ROJO):
        return _ROJO
    if any(k in texto for k in _ESTADOS_AMBAR):
        return _AMBAR
    if any(k in texto for k in _ESTADOS_VERDE):
        return _VERDE
    return None


def _marcar_estilos_aplicados(ruta):
    """openpyxl no escribe applyFont/applyFill/applyBorder y Excel puede ignorar
    rellenos y bordes sin ellos; se agregan en styles.xml del libro ya guardado."""
    import re
    import shutil
    import zipfile

    temporal = f"{ruta}.tmp"
    with zipfile.ZipFile(ruta) as origen, zipfile.ZipFile(temporal, "w", zipfile.ZIP_DEFLATED) as destino:
        for item in origen.infolist():
            datos = origen.read(item.filename)
            if item.filename == "xl/styles.xml":
                texto = datos.decode("utf-8")
                inicio, fin = texto.find("<cellXfs"), texto.find("</cellXfs>")
                if inicio != -1 and fin != -1:
                    bloque = texto[inicio:fin]
                    for atributo, bandera in (("fontId", "applyFont"), ("fillId", "applyFill"), ("borderId", "applyBorder")):
                        bloque = re.sub(
                            rf'(<xf [^>]*{atributo}="(?!0")\d+")(?![^>]*{bandera})',
                            rf'\1 {bandera}="1"', bloque,
                        )
                    texto = texto[:inicio] + bloque + texto[fin:]
                datos = texto.encode("utf-8")
            destino.writestr(item, datos)
    shutil.move(temporal, ruta)


def _estilizar_hoja_matriz(hoja, columnas, total_filas):
    """Solo apariencia: cabecera azul, filas alternadas, bordes, ajuste de texto y estados con color."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    lado = Side(style="thin", color=_BORDE)
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
    relleno_cab = PatternFill("solid", fgColor=_NAVY)
    fuente_cab = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    fuente = Font(name="Calibri", size=10, color="FF212529")

    es_estado = {i for i, c in enumerate(columnas, start=1) if str(c).lower().startswith("estado")}

    for celda in hoja[1]:
        celda.font = fuente_cab
        celda.fill = relleno_cab
        celda.border = borde
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.row_dimensions[1].height = 32

    for fila in hoja.iter_rows(min_row=2, max_row=total_filas + 1):
        zebra = PatternFill("solid", fgColor=_ZEBRA) if fila[0].row % 2 == 1 else None
        for celda in fila:
            celda.font = fuente
            celda.border = borde
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            color = _color_estado(celda.value) if celda.column in es_estado else None
            if color:
                celda.fill = PatternFill("solid", fgColor=color)
                celda.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            elif zebra:
                celda.fill = zebra

    for indice in range(1, len(columnas) + 1):
        ancho = hoja.column_dimensions[get_column_letter(indice)].width or 12
        hoja.column_dimensions[get_column_letter(indice)].width = max(12, min(ancho, 48))

    anchos = [hoja.column_dimensions[get_column_letter(i)].width for i in range(1, len(columnas) + 1)]
    for fila in hoja.iter_rows(min_row=2, max_row=total_filas + 1):
        lineas = 1
        for celda, ancho in zip(fila, anchos):
            texto = str(celda.value or "")
            partes = texto.splitlines() or [""]
            lineas = max(lineas, sum(max(1, -(-len(parte) // max(int(ancho * 1.05), 1))) for parte in partes))
        hoja.row_dimensions[fila[0].row].height = min(13.5 * lineas + 4, 300)

    hoja.sheet_view.showGridLines = False
    hoja.sheet_properties.tabColor = _NAVY
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_title_rows = "1:1"


def _estilizar_hoja_resumen(hoja):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    lado = Side(style="thin", color=_BORDE)
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
    for fila in hoja.iter_rows():
        clave, valor = fila[0], fila[1]
        clave.font = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
        clave.fill = PatternFill("solid", fgColor=_NAVY)
        valor.font = Font(name="Calibri", size=10, color="FF212529")
        valor.fill = PatternFill("solid", fgColor=_LIGHT)
        valor.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        clave.alignment = Alignment(vertical="center", wrap_text=True)
        clave.border = valor.border = borde
    hoja.column_dimensions["A"].width = 38
    hoja.column_dimensions["B"].width = 24
    hoja.sheet_view.showGridLines = False
    hoja.sheet_properties.tabColor = "FF2E75B6"


def importar_matriz_csv(archivo) -> list:
    """
    Importa una matriz de Requerimientos desde un CSV (ruta o archivo tipo
    Streamlit UploadedFile) y la normaliza al mismo contrato interno que
    produce el Mapper de GitLab (codigo, nombre, descripcion, tipo,
    historia_origen, estado). GitLab y Excel convergen al mismo formato.
    """
    from core.design_context import normalizar_fila_trazabilidad

    if hasattr(archivo, "read"):
        contenido = archivo.read()
        if isinstance(contenido, bytes):
            contenido = contenido.decode("utf-8-sig")
        lineas = contenido.splitlines()
        filas_crudas = list(csv.DictReader(lineas, delimiter=";"))
        if filas_crudas and len(filas_crudas[0]) <= 1:
            filas_crudas = list(csv.DictReader(lineas))
    else:
        with open(archivo, encoding="utf-8-sig", newline="") as handle:
            filas_crudas = list(csv.DictReader(handle, delimiter=";"))
        if filas_crudas and len(filas_crudas[0]) <= 1:
            with open(archivo, encoding="utf-8-sig", newline="") as handle:
                filas_crudas = list(csv.DictReader(handle))

    return [normalizar_fila_trazabilidad(fila) for fila in filas_crudas]


def importar_matriz_xlsx(archivo) -> list:
    """
    Importa una matriz de Requerimientos desde un XLSX (ruta o archivo tipo
    Streamlit UploadedFile) y la normaliza al mismo contrato interno que
    produce el Mapper de GitLab. GitLab y Excel convergen al mismo formato.
    """
    from openpyxl import load_workbook

    from core.design_context import normalizar_fila_trazabilidad

    workbook = load_workbook(archivo, data_only=True)
    hoja = workbook.active
    filas_iter = hoja.iter_rows(values_only=True)
    encabezados = next(filas_iter, None)
    if not encabezados:
        return []

    filas_crudas = [
        {encabezado: valor for encabezado, valor in zip(encabezados, fila_valores)}
        for fila_valores in filas_iter
        if not all(valor is None for valor in fila_valores)
    ]

    return [normalizar_fila_trazabilidad(fila) for fila in filas_crudas]
