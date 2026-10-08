"""Estilo visual compartido de los documentos generados (PDF y DOCX).

Solo define apariencia: colores, tipografía, encabezados, pies de página,
tablas y nombres de archivo. No decide ni calcula contenido.
"""
import os
import re
from datetime import datetime

NAVY = (31, 56, 100)
ACCENT = (46, 117, 182)
LIGHT = (234, 240, 248)
TEXTO = (33, 37, 41)
GRIS = (100, 106, 115)
REGLA = (191, 200, 214)

_HEX_NAVY = "1F3864"
_HEX_ACCENT = "2E75B6"
_HEX_LIGHT = "EAF0F8"
_HEX_ZEBRA = "F5F8FC"
_HEX_BORDE = "BFC8D6"

REPORTE = "Reporte_de_Ejecucion"
FORMAL = "Documento_Formal"


def nombre_archivo(tipo: str, etapa: str, extension: str, fecha: str | None = None) -> str:
    """Nombre de descarga: <tipo>_<etapa>_<fecha>.<ext> (ej. Reporte_de_Ejecucion_Diseno_2026-10-07.pdf)."""
    fecha = fecha or datetime.now().strftime("%Y-%m-%d")
    return f"{tipo}_{etapa}_{fecha}.{extension.lstrip('.')}"


# ---------------------------------------------------------------- PDF

_FUENTES_TTF = (
    (os.path.join(os.environ.get("SystemRoot", "C:/Windows"), "Fonts"),
     ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf")),
    ("/usr/share/fonts/truetype/dejavu",
     ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf")),
    ("/usr/share/fonts/truetype/liberation",
     ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf",
      "LiberationSans-Italic.ttf", "LiberationSans-BoldItalic.ttf")),
    ("/System/Library/Fonts/Supplemental",
     ("Arial.ttf", "Arial Bold.ttf", "Arial Italic.ttf", "Arial Bold Italic.ttf")),
)

_RE_ETIQUETA = re.compile(r"^([^:\n]{2,70}?):(?:\s+(.*))?$", re.S)


def _registrar_fuente(pdf) -> str | None:
    for carpeta, nombres in _FUENTES_TTF:
        rutas = [os.path.join(carpeta, n) for n in nombres]
        if all(os.path.isfile(r) for r in rutas):
            for estilo, ruta in zip(("", "B", "I", "BI"), rutas):
                pdf.add_font("Doc", estilo, ruta)
            return "Doc"
    return None


def preparar_pdf(etiqueta: str, proyecto: str, limpiar):
    """Crea el PDF con encabezado/pie y devuelve (pdf, encabezado, linea, bullet, nota).

    `limpiar` es el limpiador ASCII que solo se usa si no hay una fuente
    TrueType disponible en el sistema (con fuente TTF se conservan las tildes).
    """
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    fecha = datetime.now().strftime("%d/%m/%Y")
    estado = {"fam": "Helvetica", "unicode": False}

    def t(text):
        if not isinstance(text, str):
            text = str(text)
        return text.replace("\r", "") if estado["unicode"] else limpiar(text)

    class _PDF(FPDF):
        def header(self):
            self.set_y(10)
            self.set_font(estado["fam"], "", 8)
            self.set_text_color(*GRIS)
            mitad = (self.w - self.l_margin - self.r_margin) / 2
            self.cell(mitad, 5, t(etiqueta), new_x=XPos.RIGHT, new_y=YPos.TOP)
            self.cell(mitad, 5, t(fecha), align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            self.set_draw_color(*REGLA)
            self.set_line_width(0.3)
            self.line(self.l_margin, 16, self.w - self.r_margin, 16)
            self.set_y(self.t_margin)

        def footer(self):
            self.set_y(-14)
            self.set_draw_color(*REGLA)
            self.set_line_width(0.3)
            self.line(self.l_margin, self.get_y(), self.w - self.r_margin, self.get_y())
            self.ln(1.5)
            self.set_font(estado["fam"], "", 8)
            self.set_text_color(*GRIS)
            mitad = (self.w - self.l_margin - self.r_margin) / 2
            self.cell(mitad, 5, t(proyecto)[:70], new_x=XPos.RIGHT, new_y=YPos.TOP)
            self.cell(mitad, 5, t(f"Página {self.page_no()} de {{nb}}"), align="R")

    pdf = _PDF()
    fam = _registrar_fuente(pdf)
    if fam:
        estado.update(fam=fam, unicode=True)
    pdf.set_margins(20, 24, 20)
    pdf.set_auto_page_break(auto=True, margin=22)
    pdf.alias_nb_pages()
    pdf.add_page()
    kw = {"new_x": XPos.LMARGIN, "new_y": YPos.NEXT}
    fam = estado["fam"]

    def fuente(estilo="", size=10):
        pdf.set_font(fam, estilo, size)

    def espacio_para(alto):
        if pdf.get_y() + alto > pdf.h - pdf.b_margin:
            pdf.add_page()

    def encabezado(text, size=14):
        text = t(text)
        if size >= 16:
            pdf.set_fill_color(*NAVY)
            pdf.set_text_color(255, 255, 255)
            fuente("B", 17)
            pdf.multi_cell(0, 9, text, fill=True, padding=(3, 4, 3, 4), **kw)
            pdf.set_text_color(*TEXTO)
            pdf.ln(4)
            return
        espacio_para(28)
        if size >= 13:
            pdf.ln(3)
            x0, y0 = pdf.l_margin, pdf.get_y()
            pdf.set_fill_color(*LIGHT)
            pdf.set_text_color(*NAVY)
            fuente("B", 12.5)
            pdf.multi_cell(0, 6.5, text, fill=True, padding=(1.5, 2, 1.5, 5), **kw)
            y1 = pdf.get_y()
            if y1 > y0:
                pdf.set_fill_color(*ACCENT)
                pdf.rect(x0, y0, 1.8, y1 - y0, "F")
            pdf.ln(2)
        elif size >= 12:
            pdf.ln(3)
            pdf.set_text_color(*NAVY)
            fuente("B", 12)
            pdf.multi_cell(0, 6.5, text, **kw)
            y = pdf.get_y()
            pdf.set_draw_color(*ACCENT)
            pdf.set_line_width(0.5)
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.set_line_width(0.2)
            pdf.ln(2)
        elif size >= 11:
            pdf.ln(2)
            pdf.set_text_color(*ACCENT)
            fuente("B", 11)
            pdf.multi_cell(0, 6, text, **kw)
            pdf.ln(0.5)
        else:
            pdf.ln(1)
            pdf.set_text_color(*TEXTO)
            fuente("B", 10)
            pdf.multi_cell(0, 6, text, **kw)
        pdf.set_text_color(*TEXTO)

    def parrafo(text, sangria=0, vineta=False):
        text = t(text)
        if not text.strip():
            pdf.ln(3)
            return
        espacio_para(10)
        base = pdf.l_margin
        pdf.set_left_margin(base + sangria)
        pdf.set_x(base + sangria)
        if vineta:
            pdf.set_fill_color(*ACCENT)
            pdf.rect(base + sangria - 3.6, pdf.get_y() + 2.1, 1.4, 1.4, "F")
        alto = 5.8
        m = _RE_ETIQUETA.match(text)
        if m and "http" not in m.group(1).lower() and len(m.group(1).split()) <= 9:
            etiqueta_txt, valor = m.group(1), m.group(2) or ""
            pdf.set_text_color(*NAVY)
            fuente("B", 10)
            pdf.write(alto, etiqueta_txt + ":" + (" " if valor else ""))
            if valor:
                pdf.set_text_color(*TEXTO)
                fuente("", 10)
                pdf.write(alto, valor)
        else:
            pdf.set_text_color(*TEXTO)
            fuente("", 10)
            pdf.write(alto, text)
        pdf.ln(alto + 0.6)
        pdf.set_left_margin(base)
        pdf.set_x(base)
        pdf.set_text_color(*TEXTO)

    def linea(text):
        crudo = text if isinstance(text, str) else str(text)
        if crudo.startswith("- "):
            parrafo(crudo[2:], sangria=5, vineta=True)
        elif crudo.startswith("  "):
            parrafo(crudo.strip(), sangria=9)
        else:
            parrafo(crudo)

    def bullet(text):
        linea(f"- {text}")

    def nota(text):
        pdf.set_fill_color(*LIGHT)
        pdf.set_text_color(*GRIS)
        fuente("I", 9)
        pdf.multi_cell(0, 5, t(text), fill=True, padding=(2, 3, 2, 3), **kw)
        pdf.set_text_color(*TEXTO)
        pdf.ln(2)

    return pdf, encabezado, linea, bullet, nota


# ---------------------------------------------------------------- DOCX

_ORDEN_PPR = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
    "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]
_ORDEN_TCPR = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
    "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]
_ORDEN_TBLPR = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
    "shd", "tblLayout", "tblCellMar", "tblLook",
]


def _local(el):
    return el.tag.rsplit("}", 1)[-1]


def _insertar_ordenado(padre, nuevo, orden):
    """Inserta `nuevo` respetando el orden del esquema OOXML; reemplaza uno previo del mismo tipo."""
    nombre = _local(nuevo)
    for hijo in list(padre):
        if _local(hijo) == nombre:
            padre.remove(hijo)
    idx = orden.index(nombre)
    for hijo in padre:
        n = _local(hijo)
        if n in orden and orden.index(n) > idx:
            hijo.addprevious(nuevo)
            return
    padre.append(nuevo)


def _el(nombre, **attrs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    el = OxmlElement(f"w:{nombre}")
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


def _borde_inferior(pPr, color, sz=8, space=4):
    pbdr = _el("pBdr")
    pbdr.append(_el("bottom", val="single", sz=sz, space=space, color=color))
    _insertar_ordenado(pPr, pbdr, _ORDEN_PPR)


def _sombrear_parrafo(pPr, color, borde_izq=None):
    if borde_izq:
        pbdr = _el("pBdr")
        pbdr.append(_el("left", val="single", sz=24, space=6, color=borde_izq))
        _insertar_ordenado(pPr, pbdr, _ORDEN_PPR)
    _insertar_ordenado(pPr, _el("shd", val="clear", color="auto", fill=color), _ORDEN_PPR)


def _fuente_estilo(style, nombre, size=None, bold=None, italic=None, color=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    style.font.name = nombre
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(a), None)
    rfonts.set(qn("w:eastAsia"), nombre)
    rfonts.set(qn("w:cs"), nombre)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = RGBColor(*color)


def _campo(parrafo, instruccion, size=8, color=GRIS):
    from docx.shared import Pt, RGBColor
    run = parrafo.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(*color)
    run._r.append(_el("fldChar", fldCharType="begin"))
    instr = _el("instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = f" {instruccion} "
    run._r.append(instr)
    run._r.append(_el("fldChar", fldCharType="separate"))
    texto = _el("t")
    texto.text = "1"
    run._r.append(texto)
    run._r.append(_el("fldChar", fldCharType="end"))


def aplicar_estilo_docx(doc, etiqueta: str, proyecto: str):
    """Tipografía, encabezados, márgenes, cabecera y pie del documento Word."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    st = doc.styles
    _fuente_estilo(st["Normal"], "Calibri", size=10.5, color=TEXTO)
    st["Normal"].paragraph_format.space_after = Pt(6)
    st["Normal"].paragraph_format.line_spacing = 1.15

    _fuente_estilo(st["Title"], "Calibri", size=28, bold=True, color=NAVY)
    ppr = st["Title"].element.get_or_add_pPr()
    _borde_inferior(ppr, _HEX_ACCENT, sz=18, space=6)
    st["Title"].paragraph_format.space_after = Pt(14)

    for nombre, size, color, antes, despues in (
        ("Heading 1", 16, NAVY, 18, 6),
        ("Heading 2", 13, ACCENT, 14, 4),
        ("Heading 3", 11.5, NAVY, 10, 3),
    ):
        _fuente_estilo(st[nombre], "Calibri", size=size, bold=True, color=color)
        st[nombre].paragraph_format.space_before = Pt(antes)
        st[nombre].paragraph_format.space_after = Pt(despues)
        st[nombre].paragraph_format.keep_with_next = True
    _borde_inferior(st["Heading 1"].element.get_or_add_pPr(), _HEX_BORDE, sz=6, space=2)

    for nombre in ("List Bullet", "List Number"):
        if nombre in [s.name for s in st]:
            st[nombre].paragraph_format.space_after = Pt(3)

    seccion = doc.sections[0]
    seccion.left_margin = seccion.right_margin = Cm(2.2)
    seccion.top_margin = Cm(2.4)
    seccion.bottom_margin = Cm(2.2)

    cab = seccion.header.paragraphs[0]
    cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = cab.add_run(f"{etiqueta}  ·  {datetime.now().strftime('%d/%m/%Y')}")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(*GRIS)
    _borde_inferior(cab._p.get_or_add_pPr(), _HEX_BORDE, sz=4, space=3)

    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = pie.add_run(f"{proyecto}  ·  Página ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(*GRIS)
    _campo(pie, "PAGE")
    r = pie.add_run(" de ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(*GRIS)
    _campo(pie, "NUMPAGES")

    doc.core_properties.title = etiqueta
    doc.core_properties.author = "Modelo Multiagente"


def seccion_horizontal(doc):
    """Inicia una sección en hoja apaisada (para matrices anchas)."""
    from docx.enum.section import WD_ORIENT, WD_SECTION
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    if s.page_width < s.page_height:
        s.page_width, s.page_height = s.page_height, s.page_width
    return s


def seccion_vertical(doc):
    """Vuelve a hoja vertical tras una sección apaisada."""
    from docx.enum.section import WD_ORIENT, WD_SECTION
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.PORTRAIT
    if s.page_width > s.page_height:
        s.page_width, s.page_height = s.page_height, s.page_width
    return s


def _estilizar_tabla(table):
    from docx.shared import Pt, RGBColor

    columnas = len(table.columns)
    size = 9.5 if columnas <= 4 else 8.5 if columnas <= 7 else 7.5

    tblPr = table._tbl.tblPr
    bordes = _el("tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        bordes.append(_el(lado, val="single", sz=4, space=0, color=_HEX_BORDE))
    _insertar_ordenado(tblPr, bordes, _ORDEN_TBLPR)
    mar = _el("tblCellMar")
    for lado, w in (("top", 50), ("left", 90), ("bottom", 50), ("right", 90)):
        mar.append(_el(lado, w=w, type="dxa"))
    _insertar_ordenado(tblPr, mar, _ORDEN_TBLPR)

    for i, fila in enumerate(table.rows):
        if i == 0:
            trPr = fila._tr.get_or_add_trPr()
            trPr.append(_el("cantSplit"))
            trPr.append(_el("tblHeader"))
        for celda in fila.cells:
            if i == 0:
                relleno = _HEX_NAVY
            elif i % 2 == 0:
                relleno = _HEX_ZEBRA
            else:
                relleno = None
            if relleno:
                tcPr = celda._tc.get_or_add_tcPr()
                _insertar_ordenado(tcPr, _el("shd", val="clear", color="auto", fill=relleno), _ORDEN_TCPR)
            for p in celda.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(size)
                    if i == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def finalizar_docx(doc, notas=()):
    """Embellece tablas y párrafos "Etiqueta: valor" del documento ya armado.

    `notas`: textos (p. ej. el aviso de apoyo a la decisión) que se muestran como recuadro.
    """
    from docx.shared import Pt, RGBColor

    for table in doc.tables:
        _estilizar_tabla(table)

    notas = {n.strip() for n in notas if n}
    for p in doc.paragraphs:
        if p.style.name != "Normal" or len(p.runs) != 1:
            continue
        texto = p.runs[0].text
        if not texto.strip() or "\n" in texto or "\t" in texto:
            continue
        if texto.strip() in notas:
            run = p.runs[0]
            run.font.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(*GRIS)
            _sombrear_parrafo(p._p.get_or_add_pPr(), _HEX_LIGHT, borde_izq=_HEX_ACCENT)
            continue
        m = _RE_ETIQUETA.match(texto)
        if not m or "http" in m.group(1).lower() or len(m.group(1).split()) > 9:
            continue
        etiqueta, valor = m.group(1), m.group(2) or ""
        run = p.runs[0]
        run.text = etiqueta + ":"
        run.font.bold = True
        run.font.color.rgb = RGBColor(*NAVY)
        if valor:
            p.add_run(" " + valor)
